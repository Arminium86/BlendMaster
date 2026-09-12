from __future__ import annotations

import json
import math
import os
import re
from datetime import date, datetime

import pandas as pd


class SpreadsheetReportExportError(RuntimeError):
    """Raised when a report cannot be exported to a spreadsheet."""


class SpreadsheetReportExporter:
    """Create styled, auditable XLSX reports and lossless CSV extracts."""

    HEADER_FILL = "315B7D"
    HEADER_FONT = "FFFFFF"
    TITLE_FONT = "17365D"
    BAND_FILL = "F2F6FA"
    BORDER = "D7E0E8"

    @staticmethod
    def _frame(value):
        if isinstance(value, pd.DataFrame):
            return value.copy()
        return pd.DataFrame(value or [])

    @staticmethod
    def _sheet_name(value, existing):
        name = re.sub(r"[\\/*?:\[\]]", "_", str(value or "Report")).strip()
        name = (name or "Report")[:31]
        candidate = name
        suffix = 2
        while candidate.lower() in existing:
            ending = f" {suffix}"
            candidate = f"{name[:31 - len(ending)]}{ending}"
            suffix += 1
        existing.add(candidate.lower())
        return candidate

    @staticmethod
    def _value(value):
        if value is None or value is pd.NA:
            return None
        if isinstance(value, pd.Timestamp):
            return value.to_pydatetime()
        if isinstance(value, (datetime, date, bool, int, float, str)):
            if isinstance(value, float) and not math.isfinite(value):
                return None
            if isinstance(value, str):
                value = re.sub(
                    r"[\x00-\x08\x0B-\x0C\x0E-\x1F]", "", value
                )
                return value[:32767]
            return value
        if hasattr(value, "item"):
            try:
                return SpreadsheetReportExporter._value(value.item())
            except (TypeError, ValueError):
                pass
        try:
            if pd.isna(value):
                return None
        except (TypeError, ValueError):
            pass
        if isinstance(value, (dict, list, tuple, set)):
            return json.dumps(value, ensure_ascii=False, sort_keys=True)
        return SpreadsheetReportExporter._value(str(value))

    @staticmethod
    def _column_number_format(column):
        normalized = re.sub(
            r"[^a-z0-9]+", "_", str(column or "").lower()
        ).strip("_")
        if any(marker in normalized for marker in (
            "datetime", "timestamp", "start_time", "end_time"
        )):
            return "yyyy-mm-dd hh:mm:ss"
        if (
            "tonne" in normalized
            or normalized.endswith(("_wmt", "_dmt", "_balance"))
            or normalized in {"wmt", "dmt", "balance", "payload"}
        ):
            return "#,##0"
        if "grade" in normalized or "ratio" in normalized:
            return "0.0000"
        if any(marker in normalized for marker in (
            "ratio", "rate", "coverage", "priority"
        )):
            return "0.00"
        return None

    @classmethod
    def export_xlsx(
        cls,
        output_path,
        sheets,
        report_title="BlendMaster Report",
        report_datetime=None,
    ):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
            from openpyxl.utils import get_column_letter
        except ImportError as exc:
            raise SpreadsheetReportExportError(
                "XLSX export requires openpyxl. Install the project "
                "requirements and try again."
            ) from exc

        report_datetime = report_datetime or datetime.now()
        output_path = os.path.abspath(str(output_path))
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        workbook = Workbook()
        workbook.remove(workbook.active)
        workbook.properties.title = str(report_title)
        workbook.properties.creator = "BlendMaster"
        workbook.properties.created = report_datetime
        existing_names = set()
        thin_border = Side(style="thin", color=cls.BORDER)
        body_font = Font(name="Aptos", size=9)
        body_alignment = {align: Alignment(horizontal=align, vertical="top", wrap_text=True)
                          for align in ('left', 'right')}
        band_fill = PatternFill("solid", fgColor=cls.BAND_FILL)

        for requested_name, raw_frame in sheets:
            frame = cls._frame(raw_frame)
            if len(frame) > 1_048_571:
                raise SpreadsheetReportExportError(
                    f"'{requested_name}' has {len(frame):,} rows and exceeds "
                    "Excel's per-sheet row limit. Export that report as CSV."
                )
            if len(frame.columns) > 16_384:
                raise SpreadsheetReportExportError(
                    f"'{requested_name}' has {len(frame.columns):,} columns "
                    "and exceeds Excel's per-sheet column limit."
                )
            sheet_name = cls._sheet_name(requested_name, existing_names)
            worksheet = workbook.create_sheet(sheet_name)
            worksheet.sheet_view.showGridLines = False
            last_column = max(len(frame.columns), 1)
            last_column_letter = get_column_letter(last_column)
            worksheet.merge_cells(
                start_row=1, start_column=1,
                end_row=1, end_column=last_column,
            )
            title_cell = worksheet.cell(1, 1, str(requested_name))
            title_cell.font = Font(
                name="Aptos Display", size=16, bold=True,
                color=cls.TITLE_FONT,
            )
            title_cell.alignment = Alignment(vertical="center")
            worksheet.row_dimensions[1].height = 25
            worksheet.merge_cells(
                start_row=2, start_column=1,
                end_row=2, end_column=last_column,
            )
            report_time_cell = worksheet.cell(
                2, 1,
                "Report date and time: "
                + report_datetime.strftime("%Y-%m-%d %H:%M:%S"),
            )
            report_time_cell.font = Font(
                name="Aptos", size=9, color="5F6B78"
            )

            header_row = 4
            for column_index, column in enumerate(frame.columns, start=1):
                cell = worksheet.cell(header_row, column_index, str(column))
                cell.fill = PatternFill("solid", fgColor=cls.HEADER_FILL)
                cell.font = Font(
                    name="Aptos", size=9, bold=True,
                    color=cls.HEADER_FONT,
                )
                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
                cell.border = Border(bottom=thin_border)
            worksheet.row_dimensions[header_row].height = 30

            number_formats = [cls._column_number_format(column) for column in frame.columns]
            for row_index, row in enumerate(
                frame.itertuples(index=False, name=None), start=header_row + 1
            ):
                for column_index, value in enumerate(row, start=1):
                    column = frame.columns[column_index - 1]
                    cell = worksheet.cell(
                        row_index, column_index, cls._value(value)
                    )
                    cell.font = body_font
                    cell.alignment = body_alignment[
                            "right" if isinstance(cell.value, (int, float))
                            and not isinstance(cell.value, bool) else "left"
                        ]
                    if row_index % 2 == 0:
                        cell.fill = band_fill
                    number_format = number_formats[column_index - 1]
                    if number_format:
                        cell.number_format = number_format

            if len(frame.columns):
                worksheet.auto_filter.ref = (
                    f"A{header_row}:{last_column_letter}{header_row + len(frame)}"
                )
                worksheet.freeze_panes = f"A{header_row + 1}"
            for column_index, column in enumerate(frame.columns, start=1):
                values = [str(column)] + [
                    "" if cls._value(value) is None else str(cls._value(value))
                    for value in frame[column].head(250)
                ]
                width = min(max(max(map(len, values), default=8) + 2, 10), 45)
                worksheet.column_dimensions[
                    get_column_letter(column_index)
                ].width = width
                normalized_column = re.sub(
                    r"[^a-z0-9]+", "_", str(column).lower()
                ).strip("_")
                if any(marker in normalized_column for marker in (
                    "datetime", "timestamp", "start_time", "end_time"
                )):
                    worksheet.column_dimensions[
                        get_column_letter(column_index)
                    ].width = max(width, 21)
            worksheet.sheet_properties.pageSetUpPr.fitToPage = True
            worksheet.page_setup.orientation = "landscape"
            worksheet.page_setup.fitToWidth = 1
            worksheet.page_setup.fitToHeight = 0
            worksheet.print_title_rows = f"{header_row}:{header_row}"

        if not workbook.worksheets:
            workbook.create_sheet("Report")
        try:
            workbook.save(output_path)
        except Exception as exc:
            raise SpreadsheetReportExportError(
                f"Unable to create XLSX report: {exc}"
            ) from exc
        return output_path

    @classmethod
    def export_csv(cls, output_path, frame):
        output_path = os.path.abspath(str(output_path))
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        try:
            cls._frame(frame).to_csv(
                output_path, index=False, encoding="utf-8-sig"
            )
        except Exception as exc:
            raise SpreadsheetReportExportError(
                f"Unable to create CSV report: {exc}"
            ) from exc
        return output_path
