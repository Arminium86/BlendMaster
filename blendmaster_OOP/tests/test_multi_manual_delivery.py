"""Exercise the actual operational export buttons with distinct manual plans."""
import os
os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
from contextlib import closing
from copy import deepcopy
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch
import pandas as pd
from PyQt5.QtWidgets import QApplication, QWidget
from classes.MultiManualPlan import MultiManualPlanner
from classes.SavedPlanStore import write_manual_snapshot
from classes.MaterialFlowReview import saved_flow_data
from classes.OperationalBlendPlans import split_blend_plans
from classes.CrossFeatureReports import cross_feature_sheets
from GUI.WorkflowDependencies import manual_revision
from GUI.OperationalBlendPlanView import OperationalBlendPlanView
from GUI.OperationalPlanBackups import callbacks
from database.DatabaseContext import get_database_path, set_database_path
from tests.test_multi_manual_authoring import inputs, START


class ManualDeliveryTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls): cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.folder = tempfile.TemporaryDirectory(); self.addCleanup(self.folder.cleanup)
        self.path = Path(self.folder.name)
        previous = get_database_path(); self.addCleanup(set_database_path, previous)
        set_database_path(str(self.path/'state.db'))
        args = inputs(); planner = MultiManualPlanner(*args)
        self.report = planner.build_report(planner.build_steady_states())
        write_manual_snapshot(get_database_path(), 'Plan A', self.report)
        host = QWidget(); self.addCleanup(host.deleteLater); self.host = host
        host.active_manual_plan_id = 'Plan A'; host.manual_point_drafts = args[0]
        host.start_time_choice = START; host.planning_period_count = lambda: 3
        host.calendar_inputs = args[6]; host.product_targets = []
        host.multi_feed_configuration = args[6]['site_context']['multi_feed_settings']
        host.current_multi_feed_configuration = lambda: host.multi_feed_configuration
        host.stockpile_data = {'BACK_A': {'nearest_crusher': 'ROM'}, 'BACK_B': {'nearest_crusher': 'ROM'}}
        host.save_active_scenario_state = lambda: None
        host.manual_input_revision = manual_revision(host)
        host.manual_plan_states = {'Plan A': dict(manual_point_drafts=deepcopy(host.manual_point_drafts),
            manual_input_revision=host.manual_input_revision, blend_plan_backup_destinations={'A': 'BACK_A', 'B': 'BACK_B'})}
        host.blend_plan_backup_destinations = {'A': 'BACK_A', 'B': 'BACK_B'}
        with closing(sqlite3.connect(get_database_path())) as connection:
            pd.DataFrame([dict(plan_id='Plan A', plan_type='manual', fallback_1_destination='BACK_A', fallback_2_destination='BACK_B'),
                          dict(plan_id='Other', plan_type='manual', fallback_1_destination='WRONG')]).to_sql('material_destination_plan', connection, index=False, if_exists='replace')
        context, save = callbacks(host, 'manual')
        self.view = OperationalBlendPlanView(host, plan_type='manual', backup_context=context, save_backups=save)
        data = saved_flow_data('Plan A', get_database_path(), plan_type='manual')
        self.view.set_data(data, split_blend_plans(self.report), cross_feature_sheets('Plan A', get_database_path(), plan_type='manual'))

    def test_backup_choices_and_selections_survive_refresh_and_are_plan_scoped(self):
        self.assertEqual(self.view.backups.fields['A'].currentData(), 'BACK_A')
        self.assertEqual(set(self.view.choices['B']), {'BACK_A', 'BACK_B'})
        self.view.backups.fields['A'].setCurrentIndex(self.view.backups.fields['A'].findData('BACK_B'))
        self.assertEqual(self.host.manual_plan_states['Plan A']['blend_plan_backup_destinations']['A'], 'BACK_B')
        self.view.refresh()
        self.assertEqual(self.view.backups.fields['A'].currentData(), 'BACK_B')

    def test_selected_manual_plan_exports_after_switching_to_a_different_draft(self):
        self.host.active_manual_plan_id = 'Other'
        self.host.manual_point_drafts = {'A': dict(recipes=[], sequence=[], rates={})}
        self.host.manual_input_revision = None
        from openpyxl import load_workbook
        with patch('GUI.OperationalBlendPlanView.QMessageBox.warning') as warning:
            with patch('GUI.OperationalBlendPlanView.QFileDialog.getSaveFileName', return_value=(str(self.path/'manual.xlsx'), '')):
                self.view.export_xlsx()
            self.assertTrue((self.path/'manual.xlsx').exists(), self.view.status.text())
            with patch('GUI.OperationalBlendPlanView.QFileDialog.getSaveFileName', return_value=(str(self.path/'manual.pdf'), '')):
                self.view.export_pdf()
            self.assertTrue((self.path/'manual.pdf').exists(), self.view.status.text())
            warning.assert_not_called()
        book = load_workbook(self.path/'manual.xlsx', read_only=True)
        self.assertTrue({'A Summary', 'B Summary', 'Backup Destinations', 'Plan Readiness'}.issubset(book.sheetnames))
        book.close()
        self.assertTrue((self.path/'manual.pdf').read_bytes().startswith(b'%PDF'))
        self.assertEqual(self.host.active_manual_plan_id, 'Other')

    def test_export_reports_an_actionable_error_when_inputs_changed(self):
        self.host.manual_point_drafts['A']['rates']['Preplan'] = 5
        with patch('GUI.OperationalBlendPlanView.QMessageBox.warning') as warning:
            self.view.export_xlsx()
        warning.assert_called_once()
        self.assertIn('Inputs changed', warning.call_args.args[2])


if __name__ == '__main__': unittest.main()
