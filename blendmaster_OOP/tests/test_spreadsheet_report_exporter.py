import tempfile
import unittest
from datetime import datetime
from pathlib import Path

import pandas as pd

from classes.SpreadsheetReportExporter import SpreadsheetReportExporter


class SpreadsheetReportExporterTests(unittest.TestCase):
    def test_xlsx_contains_styled_report_sheets_and_typed_values(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "blend_plan.xlsx"
            report_time = datetime(2026, 8, 13, 9, 45, 12)
            SpreadsheetReportExporter.export_xlsx(
                output,
                [
                    ("Blend Summary", pd.DataFrame([{
                        "source": "SP1",
                        "source_actual_tonnes": 1234.6,
                        "source_grade_fe": 58.126,
                    }])),
                    ("Material Destination Plan", pd.DataFrame([{
                        "assigned_destination": "Stockpiles/SP1",
                    }])),
                ],
                report_datetime=report_time,
            )

            workbook = load_workbook(output, data_only=False)
            self.assertEqual(
                workbook.sheetnames,
                ["Blend Summary", "Material Destination Plan"],
            )
            summary = workbook["Blend Summary"]
            self.assertIn("2026-08-13 09:45:12", summary["A2"].value)
            self.assertEqual(summary["B5"].value, 1234.6)
            self.assertEqual(summary["B5"].number_format, "#,##0")
            self.assertEqual(summary["C5"].number_format, "0.00")
            self.assertEqual(summary.freeze_panes, "A5")
            self.assertEqual(summary.auto_filter.ref, "A4:C5")

    def test_csv_keeps_all_rows_and_columns(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "report.csv"
            source = pd.DataFrame({"a": [1, 2], "b": ["x", "y"]})
            SpreadsheetReportExporter.export_csv(output, source)
            restored = pd.read_csv(output)
            pd.testing.assert_frame_equal(restored, source)


if __name__ == "__main__":
    unittest.main()
